"""PRE-REGISTERED: does a person map skill names to Common Core clusters the way the language model did?

WHY. invariance.py's unit of analysis is a Common Core K-8 cluster. Every unit (XES3G5M leaf KC,
ASSISTments skill, Eedi level-3 subject, Junyi exercise) was mapped to a cluster, or NONE, by
openai/gpt-6-sol from its name alone (MAPPER), with deepseek/deepseek-v4-pro as a second map
(MAPPER2). The two models agree on 61% of units, and no person has coded any. Under CLAUDE.md a
construct a model coded may not carry a headline until it has an independent reference with at least
one human coder. The core-and-culture verdict is a candidate headline for Paper II, so this study
gives the map that reference. Written 2026-10-07, before any human code exists.

CODER. David S. Porcaro (human:david-porcaro), blind to both model maps: the coding form shows the
unit's corpus and name (and, for Chinese names, an English translation) and the 92 clusters with the
example standards the models were shown. It never shows a model's choice. The coder may use NONE, may
flag a unit "unsure", and may add a note; neither changes the scoring.

TRANSLATION. XES3G5M and Junyi name their units in Chinese, which the coder does not read. Each
sampled Chinese name is translated literally by google/gemini-3.8-flash (a third model family, neither
mapper) under TRANSLATE_PROMPT, which asks for a translation only and never mentions standards. The
translation is a model artefact between the coder and the name; the original is shown beside it, and
agreement is also reported for the English-named corpora alone so its effect is visible.

SAMPLE. 25 units from each corpus, drawn uniformly at random (seed SEED) from the units MAPPER coded
(XES3G5M 866, ASSISTments 245, Eedi 308, Junyi 1330): 100 in all. Item order is shuffled across
corpora. out/human-sample.json holds item number -> (corpus, unit id) only; names stay outside the
repository with their corpora.

PRE-REGISTRATION
  Codes: a cluster code (e.g. 4.NF.B) or NONE; blank items are excluded and counted.
  Primary statistic: Cohen's kappa, human against MAPPER, over all coded items, exact code (NONE is a
  category). Interval: 2,000 bootstrap resamples of items within corpus (stratified), percentile 95%.
  Reported beside it: human against MAPPER2; MAPPER against MAPPER2 on the same items (the
  model-model agreement the paper already states, 0.61 over all units); raw agreement; agreement at
  grade.domain (4.NF) and at domain (NF); each per corpus and for the English-named corpora
  (ASSISTments, Eedi) alone; the share of NONE in each coder.
  VERDICTS (on the primary kappa)
    MAP VALIDATED: point >= 0.60 and lower 95% bound >= 0.40.
    MAP USABLE WITH CAVEAT: point >= 0.40, not validated.
    MAP NOT VALIDATED: point < 0.40.
    UNANSWERABLE: fewer than 80 items coded.
  WHAT FOLLOWS FOR PAPER II (fixed now). VALIDATED: the core-and-culture result may be a headline,
  stated with this kappa. USABLE WITH CAVEAT: it is reported in the results with the kappa and the
  caveat, not in the abstract. NOT VALIDATED: it is a descriptive count and is not used as evidence for
  any research question. Name-only mapping noise biases invariance.py's b_other toward zero, so
  disagreement does not by itself argue the core is smaller; it argues the estimate is not yet
  trustworthy.
  WHAT IT CANNOT SAY. One human coder is a reference, not a gold standard; a second human coder would
  give the human ceiling. The sample checks the map, not the dependence statistics built on it.

SECOND HUMAN CODER (added 2026-10-07, before any codes were returned). A second person codes the same
100 items in an identical copy of the workbook, blind to both models and to the first coder; the two do
not discuss items until both files are returned. `ceiling` reports Cohen's kappa between the two humans
(same stratified bootstrap), the second coder's kappa against MAPPER and MAPPER2, and the same
grade.domain and domain agreements. Fixed now: the human-human kappa is the ceiling the primary kappa is
read against; it does not change the registered verdict, which stays on the first coder against MAPPER.
If the primary kappa is within 0.10 of the human-human kappa, the paper says the map agrees with a
person about as well as two people agree; if it is lower by more, it says so with both numbers.

SCORER CHECK (--synthetic, no names, no model call): codes planted to agree with MAPPER on a known
share of items must return that raw agreement exactly and a kappa within 0.05 of the value computed
directly. This checks the scorer only; it is a smoke check, not a detection floor.

    python3 human_codes.py --synthetic
    python3 human_codes.py sample     # writes out/human-sample.json and the coding workbook to scratch
    python3 human_codes.py score <returned workbook .xlsx or .csv>   # writes out/human-codes.json, out/human-agreement.json
    python3 human_codes.py ceiling <first coder's file> <second coder's file> <second coder id>   # out/human-ceiling.json
"""
from __future__ import annotations

import csv
import json
import random
import sys
from collections import Counter
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(HERE))
from shared.paths import scratch_root  # noqa: E402
from shared.stats import kappa_nominal  # noqa: E402

OUT = HERE / "out"
SAMPLE = OUT / "human-sample.json"
CODES = OUT / "human-codes.json"
AGREE = OUT / "human-agreement.json"
SEED = 20261007
PER_CORPUS = 25
CORPORA = ["xes3g5m", "assistments2012", "eedi", "junyi2020"]
CHINESE = {"xes3g5m", "junyi2020"}
ENGLISH = {"assistments2012", "eedi"}
TRANSLATOR = "google/gemini-3.8-flash"
CODER = "human:david-porcaro"
LABELS = {"xes3g5m": "XES3G5M (China, primary school, after-school math)",
          "assistments2012": "ASSISTments 2012-13 (US, mostly middle school)",
          "eedi": "Eedi (UK, lower secondary)",
          "junyi2020": "Junyi Academy (Taiwan, grades 1-9)"}

TRANSLATE_PROMPT = """Translate each Chinese label below into English, literally. These are names of topics on a children's mathematics learning platform; a label may be a path from general to specific, separated by ' > '. Keep the path separators. Do not explain, interpret or add anything.

{items}

Reply with JSON only: {{"t": {{"<id>": "<English>"}}}}"""


def kit_dir() -> Path:
    d = scratch_root() / "core-vs-culture-human-coding"
    d.mkdir(parents=True, exist_ok=True)
    return d


def umap() -> dict:
    return json.loads((OUT / "unit-map.json").read_text())


def draw() -> list[dict]:
    from invariance import MAPPER
    m = umap()["maps"][MAPPER]
    rng = random.Random(SEED)
    items = []
    for c in CORPORA:
        for u in rng.sample(sorted(m[c]), PER_CORPUS):
            items.append({"corpus": c, "unit": u})
    rng.shuffle(items)
    for i, it in enumerate(items, 1):
        it["item"] = i
    return items


def translate(names: dict[str, str]) -> dict[str, str]:
    from shared import llm
    if not names:
        return {}
    ids = sorted(names)
    block = "\n".join(f"t{i} | {names[x]}" for i, x in enumerate(ids))
    meta = {}
    txt = llm.call(TRANSLATOR, TRANSLATE_PROMPT.format(items=block), SEED, cache=scratch_root() / "core-vs-culture-cache",
                   max_tokens=8000, temperature=0.0, attempts=3, provider=False, meta=meta)
    s, e = txt.find("{"), txt.rfind("}")
    got = json.loads(txt[s:e + 1])["t"]
    out = {}
    for k, v in got.items():
        if k.startswith("t") and k[1:].isdigit() and int(k[1:]) < len(ids):
            out[ids[int(k[1:])]] = str(v).strip()
    missing = [x for x in ids if x not in out]
    if missing:
        raise SystemExit(f"translation missing {len(missing)} items; rerun")
    translate.meta = {k: meta.get(k) for k in ("served_model", "provider", "timestamp_utc")}
    return out


def sample() -> None:
    from invariance import clusters, unit_names
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    items = draw()
    names = unit_names()
    zh = {f"{it['corpus']}:{it['unit']}": names[it["corpus"]][it["unit"]] for it in items if it["corpus"] in CHINESE}
    en = translate(zh)
    cl = clusters()
    codes = sorted(cl) + ["NONE"]

    OUT.mkdir(exist_ok=True)
    SAMPLE.write_text(json.dumps({
        "seed": SEED, "per_corpus": PER_CORPUS, "coder": CODER, "translator": TRANSLATOR,
        "translator_served": getattr(translate, "meta", {}), "translate_prompt": TRANSLATE_PROMPT,
        "items": [{"item": it["item"], "corpus": it["corpus"], "unit": it["unit"]} for it in items]},
        indent=1) + "\n")

    wb = openpyxl.Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for line in [
        "Common Core cluster coding: 100 skill names from four math platforms",
        "",
        "For each row on the Code sheet, choose the ONE Common Core K-8 cluster whose content the skill best matches,",
        "from the dropdown in column F. The Clusters sheet lists every cluster with example standards.",
        "Choose NONE if no cluster fits (for example competition-only puzzles, high-school content, or a non-math label).",
        "",
        "Work from the name alone, as the models did: do not look the skill up on the platform.",
        "Chinese names (XES3G5M, Junyi) have a literal machine translation beside them; the original is shown too.",
        "Mark 'unsure' (column G) with Y if you are guessing between clusters; add a note (column H) if useful.",
        "Neither changes the scoring. Leave a row blank only if you cannot code it at all.",
        "",
        "Do not share this file: the skill names come from licensed datasets.",
        "When done, save it and send it back. Expect about 1.5 to 2 hours.",
    ]:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 120

    cs = wb.create_sheet("Clusters")
    cs.append(["Cluster", "Grade", "Domain", "Example standards in the cluster"])
    for k in sorted(cl):
        g, dom, _ = k.split(".")
        cs.append([k, g, dom, "; ".join(cl[k])])
    cs.append(["NONE", "", "", "No K-8 cluster fits"])
    for col, w in zip("ABCD", (10, 7, 9, 140)):
        cs.column_dimensions[col].width = w
    for c in cs[1]:
        c.font = Font(bold=True)
    cs.freeze_panes = "A2"

    ws = wb.create_sheet("Code", 1)
    ws.append(["Item", "Platform", "Skill name (original)", "English translation (machine, literal)",
               "", "Your cluster", "Unsure (Y)", "Note"])
    for it in items:
        key = f"{it['corpus']}:{it['unit']}"
        ws.append([it["item"], LABELS[it["corpus"]], names[it["corpus"]][it["unit"]],
                   en.get(key, "") if it["corpus"] in CHINESE else "(English)", "", "", "", ""])
    dv = DataValidation(type="list", formula1=f"=Clusters!$A$2:$A${len(codes) + 1}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Not a cluster", error="Pick a cluster from the list, or NONE.")
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(items) + 1}")
    for col, w in zip("ABCDEFGH", (6, 34, 50, 50, 2, 14, 10, 40)):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    hi = PatternFill("solid", fgColor="FFF2CC")
    for r in range(2, len(items) + 2):
        ws[f"F{r}"].fill = hi
        for col in "CD":
            ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    wb.active = 1
    path = kit_dir() / "common-core-coding-100.xlsx"
    wb.save(path)
    print(f"wrote {SAMPLE.relative_to(ROOT)} and {path}")


def read_returned(path: Path) -> dict[int, dict]:
    rows = []
    if path.suffix.lower() == ".csv":
        rows = list(csv.reader(open(path, newline="", encoding="utf-8-sig")))[1:]
    else:
        import openpyxl
        ws = openpyxl.load_workbook(path, data_only=True)["Code"]
        rows = [[("" if v is None else v) for v in r] for r in ws.iter_rows(min_row=2, values_only=True)]
    out = {}
    for r in rows:
        if not r or str(r[0]).strip() == "":
            continue
        code = str(r[5]).strip().upper() if len(r) > 5 else ""
        out[int(r[0])] = {"code": code,
                          "unsure": str(r[6]).strip().upper().startswith("Y") if len(r) > 6 else False,
                          "note": str(r[7]).strip() if len(r) > 7 else ""}
    return out


def level(code: str, k: int) -> str:
    return code if code == "NONE" else ".".join(code.split(".")[:k]) if k < 3 else code


def domain(code: str) -> str:
    return code if code == "NONE" else code.split(".")[1]


def agreement(h, m1, m2, corpus, rng, B=2000) -> dict:
    h, m1, m2, corpus = map(np.asarray, (h, m1, m2, corpus))

    def stats(ix):
        a, b, c = h[ix], m1[ix], m2[ix]
        return {"kappa_human_mapper": kappa_nominal(list(a), list(b)),
                "kappa_human_mapper2": kappa_nominal(list(a), list(c)),
                "kappa_mapper_mapper2": kappa_nominal(list(b), list(c)),
                "raw_human_mapper": float(np.mean(a == b)), "raw_human_mapper2": float(np.mean(a == c)),
                "raw_mapper_mapper2": float(np.mean(b == c)),
                "raw_grade_domain_human_mapper": float(np.mean([level(x, 2) == level(y, 2) for x, y in zip(a, b)])),
                "raw_domain_human_mapper": float(np.mean([domain(x) == domain(y) for x, y in zip(a, b)])),
                "none_share": {"human": float(np.mean(a == "NONE")), "mapper": float(np.mean(b == "NONE")),
                               "mapper2": float(np.mean(c == "NONE"))}}
    allix = np.arange(len(h))
    res = stats(allix)
    groups = {c: np.where(corpus == c)[0] for c in sorted(set(corpus))}
    boots = []
    for _ in range(B):
        ix = np.concatenate([rng.choice(g, len(g)) for g in groups.values()])
        boots.append(kappa_nominal(list(h[ix]), list(m1[ix])))
    res["kappa_human_mapper_ci95"] = [float(np.percentile(boots, 2.5)), float(np.percentile(boots, 97.5))]
    res["n"] = int(len(h))
    res["per_corpus"] = {c: stats(g) | {"n": int(len(g))} for c, g in groups.items()}
    eng = np.where(np.isin(corpus, sorted(ENGLISH)))[0]
    res["english_named"] = stats(eng) | {"n": int(len(eng))}
    return res


def verdict(n, k, lo) -> str:
    if n < 80:
        return "UNANSWERABLE"
    if k >= 0.60 and lo >= 0.40:
        return "MAP VALIDATED"
    return "MAP USABLE WITH CAVEAT" if k >= 0.40 else "MAP NOT VALIDATED"


def score(path: Path) -> None:
    from invariance import MAPPER, MAPPER2
    smp = json.loads(SAMPLE.read_text())
    maps = umap()["maps"]
    got = read_returned(path)
    valid = set(umap()["clusters"]) | {"NONE"}
    rows, blank, bad = [], 0, []
    for it in smp["items"]:
        g = got.get(it["item"], {"code": ""})
        if not g["code"]:
            blank += 1
            continue
        if g["code"] not in valid:
            bad.append((it["item"], g["code"]))
            continue
        rows.append((it, g))
    if bad:
        raise SystemExit(f"codes not in the cluster list: {bad}")
    h = [g["code"] for _, g in rows]
    m1 = [maps[MAPPER][it["corpus"]][it["unit"]] for it, _ in rows]
    m2 = [maps[MAPPER2].get(it["corpus"], {}).get(it["unit"], "NONE") for it, _ in rows]
    corp = [it["corpus"] for it, _ in rows]
    res = agreement(h, m1, m2, corp, np.random.default_rng(SEED))
    res["blank"] = blank
    res["unsure"] = sum(g["unsure"] for _, g in rows)
    res["verdict"] = verdict(res["n"], res["kappa_human_mapper"], res["kappa_human_mapper_ci95"][0])
    res["coder"] = CODER
    res["mappers"] = {"MAPPER": MAPPER, "MAPPER2": MAPPER2}
    CODES.write_text(json.dumps({"coder": CODER, "blind_to": [MAPPER, MAPPER2], "returned_file": path.name,
                                 "codes": [{"item": it["item"], "corpus": it["corpus"], "unit": it["unit"],
                                            "code": g["code"], "unsure": g["unsure"]} for it, g in rows]},
                                indent=1) + "\n")
    AGREE.write_text(json.dumps(res, indent=1) + "\n")
    print(json.dumps({k: res[k] for k in ("n", "kappa_human_mapper", "kappa_human_mapper_ci95",
                                           "kappa_human_mapper2", "kappa_mapper_mapper2", "verdict")}, indent=1))


def synthetic() -> None:
    rng = np.random.default_rng(SEED)
    pool = [f"{g}.{d}.A" for g in "345678" for d in ("NF", "EE", "G")] + ["NONE"]
    m1 = list(rng.choice(pool, 100))
    m2 = list(rng.choice(pool, 100))
    agree_ix = set(rng.choice(100, 70, replace=False).tolist())
    h = [m1[i] if i in agree_ix else next(p for p in rng.permutation(pool) if p != m1[i]) for i in range(100)]
    corp = [CORPORA[i % 4] for i in range(100)]
    res = agreement(h, m1, m2, corp, rng, B=200)
    direct = kappa_nominal(h, m1)
    ok = abs(res["raw_human_mapper"] - 0.70) < 1e-9 and abs(res["kappa_human_mapper"] - direct) < 0.05
    print(f"planted raw agreement 0.70 -> {res['raw_human_mapper']:.3f}; kappa {res['kappa_human_mapper']:.3f} "
          f"(direct {direct:.3f}); verdict {verdict(100, res['kappa_human_mapper'], res['kappa_human_mapper_ci95'][0])}")
    print("SCORER CHECK PASSES" if ok else "SCORER CHECK FAILS")
    sys.exit(0 if ok else 1)


def ceiling(p1: Path, p2: Path, coder2: str) -> None:
    from invariance import MAPPER, MAPPER2
    smp = json.loads(SAMPLE.read_text())
    maps = umap()["maps"]
    valid = set(umap()["clusters"]) | {"NONE"}
    g1, g2 = read_returned(p1), read_returned(p2)
    rows = [it for it in smp["items"] if g1.get(it["item"], {}).get("code") in valid
            and g2.get(it["item"], {}).get("code") in valid]
    a = [g1[it["item"]]["code"] for it in rows]
    b = [g2[it["item"]]["code"] for it in rows]
    m1 = [maps[MAPPER][it["corpus"]][it["unit"]] for it in rows]
    m2 = [maps[MAPPER2].get(it["corpus"], {}).get(it["unit"], "NONE") for it in rows]
    corp = [it["corpus"] for it in rows]
    rng = np.random.default_rng(SEED)
    hh = agreement(a, b, m1, corp, rng)          # first coder against the second
    c2 = agreement(b, m1, m2, corp, rng)         # second coder against the models
    res = {"coders": [CODER, coder2], "n": len(rows),
           "kappa_human_human": hh["kappa_human_mapper"], "kappa_human_human_ci95": hh["kappa_human_mapper_ci95"],
           "raw_human_human": hh["raw_human_mapper"],
           "raw_grade_domain_human_human": hh["raw_grade_domain_human_mapper"],
           "raw_domain_human_human": hh["raw_domain_human_mapper"],
           "second_coder_vs_models": {k: c2[k] for k in ("kappa_human_mapper", "kappa_human_mapper_ci95",
                                                         "kappa_human_mapper2", "raw_human_mapper")},
           "per_corpus_kappa_human_human": {c: v["kappa_human_mapper"] for c, v in hh["per_corpus"].items()}}
    (OUT / "human-ceiling.json").write_text(json.dumps(res, indent=1) + "\n")
    print(json.dumps(res, indent=1))


if __name__ == "__main__":
    a = sys.argv[1:]
    if a == ["--synthetic"]:
        synthetic()
    elif a == ["sample"]:
        sample()
    elif len(a) == 4 and a[0] == "ceiling":
        ceiling(Path(a[1]), Path(a[2]), a[3])
    elif len(a) == 2 and a[0] == "score":
        score(Path(a[1]))
    else:
        sys.exit(__doc__)
